from __future__ import annotations

from pathlib import Path

from constants.date_utils import express_date_to_dbf, express_month_year_label
from constants.routes import (
    INCOME_PV_TAX,
    INCOME_REPORT_FILE,
    INCOME_REPORT_FONT,
    INCOME_REPORT_FONT_SIZE,
    INCOME_REPORT_FREEZE,
    INCOME_REPORT_HEADER_ROW,
    INCOME_REPORT_HEADERS,
    INCOME_REPORT_MONEY_FORMAT,
    INCOME_REPORT_SHEET,
    INCOME_REPORT_TEXT_FORMAT,
    INCOME_REPORT_TOTAL_COLOR,
    INCOME_REPORT_WIDTHS,
    INCOME_RV_TAX,
)
from models.income_form_values import IncomeFormValues
from models.income_report_result import IncomeReportResult
from models.income_report_row import IncomeReportRow
from services.name_match_service import compact_name, tidy_name

_TEXT_COLUMNS = (3, 4)
_CENTER_COLUMNS = (1, 3, 4)
_FIRST_MONEY_COLUMN = 5


class IncomeReportService:
    @staticmethod
    def build_rows(
        invoices: list[IncomeFormValues],
        *,
        skip_shops: frozenset[str] = frozenset(),
        start_number: int = 1,
    ) -> list[IncomeReportRow]:
        shops: dict[str, dict[str, IncomeReportRow]] = {}
        for values in invoices:
            if values.kind not in (INCOME_RV_TAX, INCOME_PV_TAX):
                continue
            key = compact_name(values.company_name)
            if key in skip_shops:
                continue
            branches = shops.setdefault(key, {})
            row = branches.get(values.branch_last5)
            if row is None:
                row = IncomeReportRow(
                    shop_name=tidy_name(values.company_name),
                    tax_id=values.shop_tax_id,
                    branch=values.branch_last5,
                )
                branches[values.branch_last5] = row
            if not row.tax_id:
                row.tax_id = values.shop_tax_id
            _add_invoice(row, values)
        rows: list[IncomeReportRow] = []
        for number, branches in enumerate(shops.values(), start=start_number):
            branch_rows = [branches[branch] for branch in sorted(branches)]
            if len(branch_rows) == 1:
                branch_rows[0].number = number
                rows.append(branch_rows[0])
                continue
            rows.extend(branch_rows)
            rows.append(_total_row(branch_rows, number))
        return rows

    @staticmethod
    def write(folder: Path, start_date: str, invoices: list[IncomeFormValues]) -> IncomeReportResult:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, Side

        month, year = express_month_year_label(start_date)
        dbf_date = express_date_to_dbf(start_date)
        path = folder / INCOME_REPORT_FILE.format(month=month, year=year)
        sheet_name = INCOME_REPORT_SHEET.format(year=dbf_date[:4], month=dbf_date[4:6])
        side = Side(style="thin")
        border = Border(left=side, right=side, top=side, bottom=side)
        font = Font(name=INCOME_REPORT_FONT, size=INCOME_REPORT_FONT_SIZE)
        total_font = Font(name=INCOME_REPORT_FONT, size=INCOME_REPORT_FONT_SIZE, color=INCOME_REPORT_TOTAL_COLOR)
        center = Alignment(horizontal="center", vertical="center")

        if path.exists():
            workbook = load_workbook(path)
            sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.create_sheet(sheet_name)
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = sheet_name
        if sheet.cell(row=INCOME_REPORT_HEADER_ROW, column=1).value is None:
            for column, (title, width) in enumerate(zip(INCOME_REPORT_HEADERS, INCOME_REPORT_WIDTHS), start=1):
                cell = sheet.cell(row=INCOME_REPORT_HEADER_ROW, column=column, value=title)
                cell.font = font
                cell.border = border
                cell.alignment = center
                sheet.column_dimensions[cell.column_letter].width = width

        existing_shops, last_row, last_number = _sheet_state(sheet)
        rows = IncomeReportService.build_rows(
            invoices,
            skip_shops=frozenset(existing_shops),
            start_number=last_number + 1,
        )
        skipped_shops = _skipped_shop_names(invoices, existing_shops)

        for offset, row in enumerate(rows, start=1):
            values = [row.number, row.shop_name, row.tax_id, row.branch, *row.amounts, None]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=last_row + offset, column=column, value=value)
                cell.font = total_font if row.is_total else font
                cell.border = border
                if column in _TEXT_COLUMNS:
                    cell.number_format = INCOME_REPORT_TEXT_FORMAT
                elif column >= _FIRST_MONEY_COLUMN:
                    cell.number_format = INCOME_REPORT_MONEY_FORMAT
                if column in _CENTER_COLUMNS:
                    cell.alignment = center

        last_column = sheet.cell(row=INCOME_REPORT_HEADER_ROW, column=len(INCOME_REPORT_HEADERS)).column_letter
        sheet.auto_filter.ref = f"A{INCOME_REPORT_HEADER_ROW}:{last_column}{last_row + len(rows)}"
        sheet.freeze_panes = INCOME_REPORT_FREEZE
        workbook.save(path)
        return IncomeReportResult(path=path, skipped_shops=skipped_shops)


def _sheet_state(sheet) -> tuple[set[str], int, int]:
    shops: set[str] = set()
    last_row = INCOME_REPORT_HEADER_ROW
    last_number = 0
    for row_index, (number, name) in enumerate(
        sheet.iter_rows(min_row=INCOME_REPORT_HEADER_ROW + 1, max_col=2, values_only=True),
        start=INCOME_REPORT_HEADER_ROW + 1,
    ):
        if not name:
            continue
        shops.add(compact_name(str(name)))
        last_row = row_index
        if isinstance(number, (int, float)):
            last_number = max(last_number, int(number))
    return shops, last_row, last_number


def _skipped_shop_names(invoices: list[IncomeFormValues], existing_shops: set[str]) -> list[str]:
    names: dict[str, str] = {}
    for values in invoices:
        key = compact_name(values.company_name)
        if values.kind in (INCOME_RV_TAX, INCOME_PV_TAX) and key in existing_shops:
            names.setdefault(key, tidy_name(values.company_name))
    return list(names.values())


def _add_invoice(row: IncomeReportRow, values: IncomeFormValues) -> None:
    if values.kind == INCOME_RV_TAX:
        row.income = round(row.income + values.total_amount + values.wht_amount - values.vat_amount, 2)
        row.sale_vat = round(row.sale_vat + values.vat_amount, 2)
        row.sale_wht = round(row.sale_wht + values.wht_amount, 2)
        return
    row.royalty = round(row.royalty + values.base_amount, 2)
    row.buy_vat = round(row.buy_vat + values.vat_amount, 2)
    row.buy_wht = round(row.buy_wht + values.wht_amount, 2)


def _total_row(branch_rows: list[IncomeReportRow], number: int) -> IncomeReportRow:
    first = branch_rows[0]
    return IncomeReportRow(
        shop_name=first.shop_name,
        tax_id=next((row.tax_id for row in branch_rows if row.tax_id), ""),
        branch="",
        income=round(sum(row.income for row in branch_rows), 2),
        sale_vat=round(sum(row.sale_vat for row in branch_rows), 2),
        sale_wht=round(sum(row.sale_wht for row in branch_rows), 2),
        royalty=round(sum(row.royalty for row in branch_rows), 2),
        buy_vat=round(sum(row.buy_vat for row in branch_rows), 2),
        buy_wht=round(sum(row.buy_wht for row in branch_rows), 2),
        number=number,
        is_total=True,
    )
