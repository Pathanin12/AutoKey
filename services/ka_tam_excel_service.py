from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from models.ka_tam_row import KaTamRow
from models.ka_tam_sheet_map import KaTamSheetMap

_LEGAL_SPECIAL = ("นิติบุคคล (special)", "นิติบุคคล(special)")
_LEGAL = ("นิติบุคคล",)
_TAX_ID = ("tax id", "taxid")
_SERVICE = ("srv",)
_VAT = ("vat",)
_WT = ("wt",)
_INVOICE = ("เลขที่ใบกำกับ",)
_PERIOD = re.compile(r"(\d{4})\s+(\d{2})")


class KaTamExcelService:
    @staticmethod
    def load_rows(excel_path: Path) -> list[KaTamRow]:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            preview: list[tuple] = []
            rows: list[tuple] = []
            for index, raw in enumerate(sheet.iter_rows(values_only=True)):
                values = tuple(raw)
                rows.append(values)
                if index < 8:
                    preview.append(values)
            column_map = detect_column_map(preview)
            period_text = f"{sheet.title} {excel_path.stem}"
            parsed: list[KaTamRow] = []
            for excel_row, raw_values in enumerate(rows[column_map.data_start_row :], start=column_map.data_start_row + 1):
                row = _parse_row(raw_values, column_map, excel_row, sheet.title, period_text)
                if row is not None:
                    parsed.append(row)
            return parsed
        finally:
            workbook.close()


def detect_column_map(preview_rows: list[tuple]) -> KaTamSheetMap:
    best: tuple[int, int, list[str]] | None = None
    for index, raw in enumerate(preview_rows[:8]):
        headers = [_cell_text(value) for value in raw]
        score = 0
        if _find_header(headers, _LEGAL_SPECIAL) is not None or _find_header(headers, _LEGAL) is not None:
            score += 2
        for group in (_SERVICE, _VAT, _WT, _INVOICE, _TAX_ID):
            if _find_header(headers, group) is not None:
                score += 1
        if best is None or score > best[0]:
            best = (score, index, headers)
    if best is None or best[0] < 4:
        raise ValueError("ไม่พบหัวคอลัมน์ นิติบุคคล / srv / vat / wt ในไฟล์ Excel")
    header_row = best[1]
    headers = best[2]
    legal_name = _find_header(headers, _LEGAL_SPECIAL) or _find_header(headers, _LEGAL)
    service_amount = _find_header(headers, _SERVICE)
    vat_amount = _find_header(headers, _VAT)
    wt_amount = _find_header(headers, _WT)
    missing = [
        label
        for label, index in (
            ("นิติบุคคล", legal_name),
            ("srv", service_amount),
            ("vat", vat_amount),
            ("wt", wt_amount),
        )
        if index is None
    ]
    if missing:
        raise ValueError("ไม่พบคอลัมน์: " + ", ".join(missing))
    assert legal_name is not None and service_amount is not None
    assert vat_amount is not None and wt_amount is not None
    data_start_row = header_row + 1
    while data_start_row < len(preview_rows) and not _is_sequence_cell(
        preview_rows[data_start_row][0] if preview_rows[data_start_row] else None
    ):
        data_start_row += 1
    return KaTamSheetMap(
        header_row=header_row,
        data_start_row=data_start_row,
        legal_name=legal_name,
        service_amount=service_amount,
        vat_amount=vat_amount,
        wt_amount=wt_amount,
        tax_id=_find_header(headers, _TAX_ID) if _find_header(headers, _TAX_ID) is not None else -1,
        invoice_number=_find_header(headers, _INVOICE) if _find_header(headers, _INVOICE) is not None else -1,
    )


def _parse_row(
    raw_values: tuple,
    column_map: KaTamSheetMap,
    excel_row_number: int,
    sheet_name: str,
    period_text: str,
) -> KaTamRow | None:
    sequence = _parse_sequence(raw_values)
    if sequence is None:
        return None
    legal_name = _to_text(_cell_at(raw_values, column_map.legal_name))
    if not legal_name:
        return None
    invoice_number = _to_text(_cell_at(raw_values, column_map.invoice_number))
    if not invoice_number:
        invoice_number = _nrg_reference(period_text, sequence)
    elif invoice_number.upper().startswith("NRG"):
        invoice_number = invoice_number
    return KaTamRow(
        row_number=excel_row_number,
        sequence=sequence,
        sheet_name=sheet_name,
        legal_name=legal_name,
        service_amount=_to_float(_cell_at(raw_values, column_map.service_amount)),
        vat_amount=_to_float(_cell_at(raw_values, column_map.vat_amount)),
        wt_amount=_to_float(_cell_at(raw_values, column_map.wt_amount)),
        invoice_number=invoice_number,
        tax_id=_to_tax_id(_cell_at(raw_values, column_map.tax_id)),
    )


def _nrg_reference(period_text: str, sequence: int) -> str:
    match = _PERIOD.search(period_text.strip())
    if not match:
        return ""
    return f"NRG{int(match.group(1))}{int(match.group(2)):02d}{int(sequence):04d}"


def _cell_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    return " ".join(text.lower().split())


def _find_header(headers: list[str], names: tuple[str, ...]) -> int | None:
    for name in names:
        if name in headers:
            return headers.index(name)
    return None


def _is_sequence_cell(value) -> bool:
    return _parse_sequence((value,)) is not None


def _parse_sequence(raw_values: tuple) -> int | None:
    if not raw_values:
        return None
    first = raw_values[0]
    if first is None:
        return None
    try:
        return int(float(first))
    except (TypeError, ValueError):
        return None


def _cell_at(raw_values: tuple, index: int):
    if index < 0 or index >= len(raw_values):
        return None
    return raw_values[index]


def _to_float(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, str):
        value = value.strip().replace(",", "")
        if not value:
            return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _to_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"xx", "nan", "no", "acct"}:
        return ""
    return text


def _to_tax_id(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return "".join(ch for ch in f"{value:.0f}" if ch.isdigit())
    return _to_text(value)
