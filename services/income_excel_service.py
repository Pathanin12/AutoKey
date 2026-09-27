from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from models.income_excel_sheet_map import IncomeExcelSheetMap
from models.income_portal_account import IncomePortalAccount

_LEGAL = ("นิติบุคคล",)
_USER = ("user",)
_PASS = ("pass",)


class IncomeExcelService:
    @staticmethod
    def load_accounts(excel_path: Path) -> list[IncomePortalAccount]:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            column_map = detect_sheet_map(workbook)
            sheet = workbook[column_map.sheet_name]
            accounts: list[IncomePortalAccount] = []
            seen: set[str] = set()
            for index, raw in enumerate(sheet.iter_rows(values_only=True)):
                if index < column_map.data_start_row:
                    continue
                account = _parse_row(raw, column_map)
                if account is None:
                    continue
                key = account.username.lower()
                if key in seen:
                    continue
                seen.add(key)
                accounts.append(account)
            return accounts
        finally:
            workbook.close()


def detect_sheet_map(workbook) -> IncomeExcelSheetMap:
    best: IncomeExcelSheetMap | None = None
    for name in workbook.sheetnames:
        sheet = workbook[name]
        preview: list[tuple] = []
        for index, raw in enumerate(sheet.iter_rows(values_only=True)):
            preview.append(tuple(raw))
            if index >= 7:
                break
        found = _map_from_preview(preview, name)
        if found is not None:
            best = found
            break
    if best is None:
        raise ValueError("ไม่พบหัวคอลัมน์ นิติบุคคล / user / pass ในไฟล์ Excel")
    return best


def _map_from_preview(preview: list[tuple], sheet_name: str) -> IncomeExcelSheetMap | None:
    for index, raw in enumerate(preview[:8]):
        headers = [_cell_text(value) for value in raw]
        legal_name = _find_header(headers, _LEGAL)
        username = _find_header(headers, _USER)
        password = _find_header(headers, _PASS)
        if legal_name is None or username is None or password is None:
            continue
        return IncomeExcelSheetMap(
            header_row=index,
            data_start_row=index + 1,
            legal_name=legal_name,
            username=username,
            password=password,
            sheet_name=sheet_name,
        )
    return None


def _parse_row(raw_values: tuple, column_map: IncomeExcelSheetMap) -> IncomePortalAccount | None:
    legal_name = _to_text(_cell_at(raw_values, column_map.legal_name))
    username = _to_text(_cell_at(raw_values, column_map.username))
    password = _to_text(_cell_at(raw_values, column_map.password))
    if not legal_name or not username or not password:
        return None
    return IncomePortalAccount(legal_name=legal_name, username=username, password=password)


def _cell_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    return " ".join(text.lower().split())


def _find_header(headers: list[str], names: tuple[str, ...]) -> int | None:
    for name in names:
        if name in headers:
            return headers.index(name)
    return None


def _cell_at(raw_values: tuple, index: int):
    if index < 0 or index >= len(raw_values):
        return None
    return raw_values[index]


def _to_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none"}:
        return ""
    return text
