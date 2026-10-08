from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from constants.routes import UI_TEXT
from models.month_year_period import MonthYearPeriod
from models.payroll_row import PayrollRow
from services.name_match_service import tidy_name

_BLOCK_WIDTH = 12
_SHOP_CODE_RE = re.compile(r"\s+\d{3,}$")
_MONTH_HEADER = "เดือน"
_NAME_HEADER = "นิติบุคคล"
_SALARY_HEADER = "เงินเดือน"
_SSO5_HEADER = "sso 5%"
_WELFARE025_HEADER = "สงเคราะห์ลูกจ้าง 0.25%"
_SSO10_HEADER = "sso 10%"
_WELFARE_HEADER = "สงเคราะห์ลูกจ้าง 0.50%"
_TAX_HEADER = "ภาษี"
_NET_HEADER = "เงินสุทธิ"


class PayrollExcelService:
    @staticmethod
    def count_rows(excel_path: Path) -> int:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            rows = [tuple(raw) for raw in workbook[workbook.sheetnames[0]].iter_rows(values_only=True)]
        finally:
            workbook.close()
        header_row, name_col, _month_cols = _layout(rows)
        return sum(1 for raw in rows[header_row + 1 :] if _to_text(_cell_at(raw, name_col)))

    @staticmethod
    def load_rows(excel_path: Path, period: MonthYearPeriod) -> list[PayrollRow]:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = [tuple(raw) for raw in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        header_row, name_col, month_cols = _layout(rows)
        headers = [_cell_text(value) for value in rows[header_row]]
        month_col = _month_col_for_period(rows[header_row + 1 :], month_cols, period)
        if month_col is None:
            raise ValueError(UI_TEXT["payroll_period_missing"].format(period=period.text))
        parsed: list[PayrollRow] = []
        for excel_row, raw in enumerate(rows[header_row + 1 :], start=header_row + 2):
            name = _to_text(_cell_at(raw, name_col))
            if not name:
                continue
            prefix = _to_text(_cell_at(raw, name_col - 1))
            parsed.append(
                PayrollRow(
                    row_number=excel_row,
                    legal_name=name,
                    match_names=_match_names(prefix, name),
                    salary=_block_amount(raw, headers, month_col, _SALARY_HEADER),
                    sso5=_block_amount(raw, headers, month_col, _SSO5_HEADER),
                    welfare025=_block_amount(raw, headers, month_col, _WELFARE025_HEADER),
                    sso10=_block_amount(raw, headers, month_col, _SSO10_HEADER),
                    welfare=_block_amount(raw, headers, month_col, _WELFARE_HEADER),
                    tax=_block_amount(raw, headers, month_col, _TAX_HEADER),
                    net=_block_amount(raw, headers, month_col, _NET_HEADER),
                )
            )
        return parsed


def _layout(rows: list[tuple]) -> tuple[int, int, list[int]]:
    for index, raw in enumerate(rows[:8]):
        headers = [_cell_text(value) for value in raw]
        name_col = _find_header(headers, (_NAME_HEADER,))
        month_cols = [i for i, header in enumerate(headers) if header == _MONTH_HEADER]
        if name_col is not None and month_cols:
            return index, name_col, month_cols
    raise ValueError("ไม่พบหัวคอลัมน์ นิติบุคคล / เดือน ในไฟล์ Excel")


def _month_col_for_period(data_rows: list[tuple], month_cols: list[int], period: MonthYearPeriod) -> int | None:
    for col in month_cols:
        for raw in data_rows:
            found = MonthYearPeriod.parse(_to_text(_cell_at(raw, col)))
            if found.is_valid and found.month == period.month and found.year == period.year:
                return col
    return None


def _block_amount(raw: tuple, headers: list[str], month_col: int, label: str) -> float:
    end = min(month_col + _BLOCK_WIDTH, len(headers))
    for index in range(month_col, end):
        header = headers[index] if index < len(headers) else ""
        if header == label or header.startswith(label):
            return _to_float(_cell_at(raw, index))
    return 0.0


def _match_names(prefix: str, name: str) -> tuple[str, ...]:
    names: list[str] = []
    for raw in (name, f"{prefix} {name}".strip(), f"{prefix}{name}".strip()):
        text = tidy_name(raw)
        if text and text not in names:
            names.append(text)
        stripped = _SHOP_CODE_RE.sub("", text).strip()
        if stripped and stripped not in names:
            names.append(stripped)
    return tuple(names)


def _cell_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    return " ".join(text.lower().split())


def _find_header(headers: list[str], names: tuple[str, ...]) -> int | None:
    for name in names:
        needle = name.lower()
        for index, header in enumerate(headers):
            if header == needle:
                return index
    return None


def _cell_at(raw_values: tuple, index: int):
    if index < 0 or index >= len(raw_values):
        return None
    return raw_values[index]


def _to_float(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, str):
        value = value.strip().replace(",", "")
        if not value or value.startswith("#"):
            return 0.0
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return 0.0


def _to_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan"}:
        return ""
    return text
