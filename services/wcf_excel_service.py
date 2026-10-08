from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from constants.date_utils import format_express_pv_date, is_complete_express_date
from models.month_year_period import MonthYearPeriod
from models.wcf_row import WcfRow
from models.wcf_sheet_map import WcfSheetMap
from services.name_match_service import tidy_name

_SHOP_CODE_RE = re.compile(r"\s*\(\d+\)\s*$")
_TIME_SPLIT_RE = re.compile(r"(?:เวลา|[T\s])+[0-9]{1,2}:[0-9]{2}")
_SHOP_HEADERS = ("ชื่อสถานประกอบการ",)
_ALT_HEADERS = ("ชื่อในระบบ",)
_DATE_HEADERS = ("วันที่ชำระ",)
_CONTRIB_HEADERS = ("เงินสมทบ",)
_SURCHARGE_HEADERS = ("เงินเพิ่ม",)
_PAID_HEADERS = ("รวมยอด",)
_YEAR_HEADERS = ("งวดที่ชำระ",)
_RECEIPT_HEADERS = ("เลขที่ใบเสร็จรับเงิน",)
_NOTICE_HEADERS = ("เลขใบแจ้ง",)


class WcfExcelService:
    @staticmethod
    def load_rows(excel_path: Path) -> list[WcfRow]:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = [tuple(raw) for raw in sheet.iter_rows(values_only=True)]
            column_map = detect_column_map(rows[:8])
            parsed: list[WcfRow] = []
            for excel_row, raw_values in enumerate(
                rows[column_map.data_start_row :], start=column_map.data_start_row + 1
            ):
                row = _parse_row(raw_values, column_map, excel_row)
                if row is not None:
                    parsed.append(row)
            return parsed
        finally:
            workbook.close()


def detect_column_map(preview_rows: list[tuple]) -> WcfSheetMap:
    best: tuple[int, int, list[str]] | None = None
    for index, raw in enumerate(preview_rows[:8]):
        headers = [_cell_text(value) for value in raw]
        score = 0
        for group in (
            _SHOP_HEADERS,
            _ALT_HEADERS,
            _DATE_HEADERS,
            _CONTRIB_HEADERS,
            _SURCHARGE_HEADERS,
            _PAID_HEADERS,
            _YEAR_HEADERS,
        ):
            if _find_header(headers, group) is not None:
                score += 1
        if best is None or score > best[0]:
            best = (score, index, headers)
    if best is None or best[0] < 4:
        raise ValueError("ไม่พบหัวคอลัมน์ ชื่อสถานประกอบการ / วันที่ชำระ / เงินสมทบ ในไฟล์ Excel")
    header_row = best[1]
    headers = best[2]
    shop_name = _find_header(headers, _SHOP_HEADERS)
    alt_name = _find_header(headers, _ALT_HEADERS)
    pay_date = _find_header(headers, _DATE_HEADERS)
    contrib_amount = _find_header(headers, _CONTRIB_HEADERS)
    missing = [
        label
        for label, index in (
            ("ชื่อสถานประกอบการ", shop_name if shop_name is not None else alt_name),
            ("วันที่ชำระ", pay_date),
            ("เงินสมทบ", contrib_amount),
        )
        if index is None
    ]
    if missing:
        raise ValueError("ไม่พบคอลัมน์: " + ", ".join(missing))
    assert pay_date is not None and contrib_amount is not None
    return WcfSheetMap(
        header_row=header_row,
        data_start_row=header_row + 1,
        shop_name=shop_name if shop_name is not None else -1,
        alt_name=alt_name if alt_name is not None else -1,
        pay_date=pay_date,
        contrib_amount=contrib_amount,
        surcharge=_find_header(headers, _SURCHARGE_HEADERS) if _find_header(headers, _SURCHARGE_HEADERS) is not None else -1,
        paid_amount=_find_header(headers, _PAID_HEADERS) if _find_header(headers, _PAID_HEADERS) is not None else -1,
        year=_find_header(headers, _YEAR_HEADERS) if _find_header(headers, _YEAR_HEADERS) is not None else -1,
        receipt_no=_find_header(headers, _RECEIPT_HEADERS) if _find_header(headers, _RECEIPT_HEADERS) is not None else -1,
        notice_no=_find_header(headers, _NOTICE_HEADERS) if _find_header(headers, _NOTICE_HEADERS) is not None else -1,
    )


def match_name(value: str) -> str:
    return _SHOP_CODE_RE.sub("", tidy_name(value)).strip()


def _parse_row(raw_values: tuple, column_map: WcfSheetMap, excel_row_number: int) -> WcfRow | None:
    shop_name = _to_text(_cell_at(raw_values, column_map.shop_name)) or _to_text(
        _cell_at(raw_values, column_map.alt_name)
    )
    if not shop_name:
        return None
    contrib = _to_float(_cell_at(raw_values, column_map.contrib_amount))
    surcharge = _to_float(_cell_at(raw_values, column_map.surcharge))
    paid = _to_float(_cell_at(raw_values, column_map.paid_amount))
    if paid <= 0:
        paid = round(contrib + surcharge, 2)
    receipt_no = _to_id(_cell_at(raw_values, column_map.receipt_no)) or _to_id(
        _cell_at(raw_values, column_map.notice_no)
    )
    pv_date = parse_wcf_pay_date(_cell_at(raw_values, column_map.pay_date))
    return WcfRow(
        row_number=excel_row_number,
        shop_name=shop_name,
        match_name=match_name(shop_name),
        pv_date=pv_date,
        contrib_amount=contrib,
        surcharge=surcharge,
        paid_amount=paid,
        period=_to_period(_cell_at(raw_values, column_map.year), pv_date),
        receipt_no=receipt_no,
    )


def parse_wcf_pay_date(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        year = value.year + 543 if value.year < 2400 else value.year
        text = f"{value.day:02d}/{value.month:02d}/{year}"
        formatted = format_express_pv_date(text)
        return formatted if is_complete_express_date(formatted) else ""
    text = _TIME_SPLIT_RE.split(str(value).strip(), maxsplit=1)[0].strip()
    formatted = format_express_pv_date(text)
    return formatted if is_complete_express_date(formatted) else ""


def _to_period(value, pv_date: str) -> str:
    period = MonthYearPeriod.parse(_to_text(value))
    if period.is_valid:
        return period.text
    if len(pv_date) >= 8:
        from_date = MonthYearPeriod.parse(pv_date[3:])
        if from_date.is_valid:
            return from_date.text
    return ""


def _cell_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    return " ".join(text.lower().split())


def _find_header(headers: list[str], names: tuple[str, ...]) -> int | None:
    for name in names:
        needle = name.lower()
        for index, header in enumerate(headers):
            if needle == header or needle in header:
                return index
    return None


def _cell_at(raw_values: tuple, index: int):
    if index < 0 or index >= len(raw_values):
        return None
    return raw_values[index]


def _to_float(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, str):
        value = value.strip().replace(",", "")
        if not value:
            return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _to_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan"}:
        return ""
    return text


def _to_id(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return "".join(ch for ch in f"{value:.0f}" if ch.isdigit())
    return _to_text(value).strip().upper()
